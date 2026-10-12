"""بوربوينت «لمسة تراث WRO» (العربي والإنجليزي) يعمل وحده بصوت الراوي
- كل شريحة بصوت = فيديو بعرض الشريحة (لقطات خطواتها + صوت الراوي، كل عنصر يظهر مع جملته)
  ينسخ شريحة «القصة كاملة» من بوربوينت عرض المحكّمين، لأن تشغيل الفيديو تلقائيًا فيها كتبه PowerPoint نفسه
- شريحة الفريق صورة بلا صوت (يتكلم الطالبان بصوتهما) وتنتقل بالنقر
- بقية الشرائح تنتقل وحدها بعد انتهاء صوتها
التشغيل (من جذر المستودع، بعد narration-tts.py وبعد story-video.js للعرض):
  python3 masmak-wro/build-pptx.py masmak-wro       ← «لمسة تراث WRO.pptx»
  python3 masmak-wro/build-pptx.py masmak-wro-en    ← «Touch of Heritage WRO.pptx»
يحتاج: python-pptx و ffmpeg ونسخة من puppeteer-core (NODE_PATH) لتصوير الخطوات (shots.js)
"""
import copy, json, os, re, subprocess, sys, tempfile
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

DECK = sys.argv[1].rstrip('/')
EN = DECK.endswith('-en')
SRC = 'masmak-judges-en/Touch of Heritage.pptx' if EN else 'masmak-judges/المصمك أمام لجنة التحكيم.pptx'
OUT = f'{DECK}/' + ('Touch of Heritage WRO.pptx' if EN else 'لمسة تراث WRO.pptx')
TAIL = 0.8                               # ثانية بعد انتهاء الصوت قبل الانتقال
HERE = os.path.dirname(os.path.abspath(__file__))
CUES = json.loads(re.search(r'=\s*(\{.*\})\s*;', open(f'{DECK}/audio/cues.js', encoding='utf-8').read(), re.S).group(1))
work = tempfile.mkdtemp(prefix='wro_pptx_')

# ١) لقطات كل شريحة بكل خطواتها + بيانات الشرائح (النوع، الصوت، الملاحظات)
subprocess.run(['node', os.path.join(HERE, 'shots.js'), DECK, work], check=True)
meta = json.load(open(f'{work}/meta.json', encoding='utf-8'))


def ffdur(f):
    return float(subprocess.check_output(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0', f]))


def make_video(i, m):
    """فيديو الشريحة: كل خطوة تظهر عند بداية جملتها، والصوت فوقها"""
    key = m['say']; audio = f'{DECK}/audio/{key}.mp3'
    total = ffdur(audio) + TAIL
    starts = CUES[key]['starts']
    n = m['steps']
    # الخطوة k تبدأ مع الجملة k (الجملة الأولى = الشريحة بلا عناصر إضافية)
    times = [0.0] + [starts[k] if k < len(starts) else starts[-1] for k in range(1, n + 1)]
    lst = os.path.join(work, f'v{i}.txt')
    with open(lst, 'w') as fh:
        for k in range(n + 1):
            end = times[k + 1] if k < n else total
            fh.write(f"file '{work}/s{i}-{k}.png'\nduration {max(end - times[k], 0.04):.3f}\n")
        fh.write(f"file '{work}/s{i}-{n}.png'\n")
    out = os.path.join(work, f'v{i}.mp4')
    subprocess.run(['ffmpeg', '-v', 'error', '-y', '-f', 'concat', '-safe', '0', '-i', lst, '-i', audio,
                    '-vf', 'fps=12,format=yuv420p', '-c:v', 'libx264', '-crf', '22', '-preset', 'medium',
                    '-c:a', 'aac', '-b:a', '128k', '-t', f'{total:.2f}', '-movflags', '+faststart', out], check=True)
    return out, total


p = Presentation(SRC)
slides = list(p.slides)
story = next(s for s in slides if any(sh.name.startswith('video-') for sh in s.shapes))
team = next(s for s in slides if s.has_notes_slide and ('يعرّف كل طالب' in s.notes_slide.notes_text_frame.text or 'Each student introduces' in s.notes_slide.notes_text_frame.text))
layout = story.slide_layout


def set_notes(slide, text, ref=story):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        src = next(sh for sh in ref.notes_slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        ns.shapes._spTree.append(copy.deepcopy(src._element))
        tf = [sh for sh in ns.shapes if sh.has_text_frame][-1].text_frame
    tf.text = text


def video_slide(mp4, poster, notes):
    """نسخة من شريحة القصة (فيديو بعرض الشريحة يبدأ وحده) بفيديو وغلاف جديدين"""
    s = p.slides.add_slide(layout)
    sp = s.shapes._spTree
    for el in list(sp)[2:]:
        sp.remove(el)
    for el in list(story.shapes._spTree)[2:]:
        sp.append(copy.deepcopy(el))
    s._element.append(copy.deepcopy(story._element.find(qn('p:timing'))))
    # العلاقات: الفيديو والوسائط والغلاف بمعرّفات جديدة
    # جزء وسائط جديد مباشرة (أجزاء الفيديو الموجودة تُقرأ كأجزاء عامة فلا تصلح للبحث بالبصمة)
    from pptx.media import Video
    from pptx.parts.media import MediaPart
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    mp = MediaPart.new(p.part.package, Video.from_path_or_file_like(mp4, 'video/mp4'))
    media_rId, video_rId = s.part.relate_to(mp, RT.MEDIA), s.part.relate_to(mp, RT.VIDEO)
    _, img_rId = s.part.get_or_add_image_part(poster)
    old = {r.reltype: rid for rid, r in story.part.rels.items()}
    rmap = {old['http://schemas.microsoft.com/office/2007/relationships/media']: media_rId,
            old['http://schemas.openxmlformats.org/officeDocument/2006/relationships/video']: video_rId,
            old['http://schemas.openxmlformats.org/officeDocument/2006/relationships/image']: img_rId}
    for el in sp.iter():
        for a, v in list(el.attrib.items()):
            if v in rmap and a.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                el.set(a, rmap[v])
    set_notes(s, notes)
    return s


def auto_advance(slide, secs):
    """انتقال تلقائي بعد مدة الصوت (والنقر يعمل أيضًا)"""
    el = slide._element
    for t in el.findall(qn('p:transition')):
        el.remove(t)
    tr = etree.SubElement(el, qn('p:transition'))
    tr.set('spd', 'med'); tr.set('advClick', '1'); tr.set('advTm', str(int(secs * 1000)))
    etree.SubElement(tr, qn('p:fade'))
    timing = el.find(qn('p:timing'))
    if timing is not None:
        timing.addprevious(tr)


new = []
for i, m in enumerate(meta):
    if m['t'] == 'team':
        new.append(team); continue
    if m['t'] == 'fullstory':
        # فيديو القصة الجديد (بالترقيم الصحيح) والغلاف من هذا العرض
        for rid, r in story.part.rels.items():
            if r.reltype.endswith('/video') or r.reltype.endswith('/media'):
                r.target_part._blob = open(f'{DECK}/story.mp4', 'rb').read()
            if r.reltype.endswith('/image'):
                # الغلاف بصيغة الجزء نفسه (jpg) حتى لا يعترض PowerPoint
                from PIL import Image; import io
                buf = io.BytesIO(); Image.open(f'{work}/s{i}-{m["steps"]}.png').convert('RGB').save(buf, 'JPEG', quality=85)
                r.target_part._blob = buf.getvalue()
        auto_advance(story, ffdur(f'{DECK}/story.mp4') + TAIL)
        new.append(story); continue
    mp4, secs = make_video(i, m)
    s = video_slide(mp4, f'{work}/s{i}-0.png', m['notes'])
    auto_advance(s, secs)
    new.append(s)

# الترتيب الجديد، وحذف بقية شرائح العرض الطويل
lst = p.slides._sldIdLst
by_id = {el.get('id'): el for el in lst}
keep = [by_id[str(s.slide_id)] for s in new]
for el in list(lst):
    if el not in keep:
        p.part.drop_rel(el.get(qn('r:id')))
    lst.remove(el)
for el in keep:
    lst.append(el)
p.core_properties.title = ('Touch of Heritage WRO' if EN else 'لمسة تراث WRO')
p.save(OUT)
print('✓', OUT, len(p.slides), 'شرائح', round(os.path.getsize(OUT) / 1e6, 1), 'MB')
